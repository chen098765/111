from docx import Document

# 创建参考答案文档
ref_doc = Document()
ref_doc.add_heading('实训目的', level=1)
ref_doc.add_paragraph('''
本实训旨在帮助学生掌握Python编程的基本技能，包括数据结构、算法设计和文件操作。通过本次实训，学生将能够：
1. 熟练使用Python基本数据类型和操作
2. 掌握列表、字典等数据结构的使用
3. 理解并实现常用算法
4. 学会文件读写操作
''')

ref_doc.add_heading('实训步骤', level=1)
ref_doc.add_paragraph('''
1. 环境搭建
   - 安装Python 3.8或更高版本
   - 配置开发环境（如PyCharm或VSCode）
   - 安装必要的第三方库

2. 数据结构练习
   - 实现列表的基本操作（增删改查）
   - 使用字典进行数据存储和查询
   - 练习集合和元组的使用

3. 算法实现
   - 实现冒泡排序算法
   - 实现二分查找算法
   - 分析算法的时间复杂度

4. 文件操作
   - 读取文本文件内容
   - 写入数据到文件
   - 处理CSV格式的数据文件
''')

ref_doc.add_heading('实验结果', level=1)
ref_doc.add_paragraph('''
通过本次实训，成功完成了以下任务：
1. 实现了冒泡排序算法，对100个随机数进行排序，用时0.05秒
2. 实现了二分查找算法，在有序数组中查找指定元素，查找效率为O(log n)
3. 成功读取和处理了包含1000条记录的CSV文件
4. 生成了包含统计信息的输出文件
''')

ref_doc.add_heading('问题反思', level=1)
ref_doc.add_paragraph('''
在实训过程中遇到了以下问题：
1. 冒泡排序效率较低，对于大数据量排序耗时较长
2. 文件编码问题导致读取中文内容时出现乱码
3. 内存占用问题在处理大文件时需要优化

通过查阅资料和调试，这些问题都得到了解决。特别是使用了更高效的排序算法（如快速排序）来替代冒泡排序，大大提高了性能。
''')

ref_doc.add_heading('心得体会', level=1)
ref_doc.add_paragraph('''
通过本次实训，我深刻体会到：
1. 理论知识需要通过实践来巩固
2. 编程能力需要通过大量的练习来提升
3. 遇到问题时要善于查找资料和寻求帮助
4. 代码规范和注释的重要性
5. 调试技巧对于解决问题至关重要

这次实训不仅提高了我的编程技能，也培养了我的问题解决能力和自主学习能力。
''')

ref_doc.save('reference_report.docx')

# 创建学生报告文档
stu_doc = Document()
stu_doc.add_heading('实训目的', level=1)
stu_doc.add_paragraph('''
本实训目的是学习Python编程基础，包括数据结构和算法。通过实训，学生可以：
1. 学习Python基本语法
2. 掌握列表和字典的使用
3. 了解常用算法
4. 学会文件操作
''')

stu_doc.add_heading('实训步骤', level=1)
stu_doc.add_paragraph('''
1. 安装Python环境
2. 练习数据结构
   - 列表操作
   - 字典使用
3. 实现算法
   - 排序算法
   - 查找算法
4. 文件处理
   - 读取文件
   - 写入文件
''')

stu_doc.add_heading('实验结果', level=1)
stu_doc.add_paragraph('''
完成了以下任务：
1. 实现了排序算法
2. 实现了查找算法
3. 处理了CSV文件
4. 生成了输出文件
''')

stu_doc.add_heading('问题反思', level=1)
stu_doc.add_paragraph('''
遇到的问题：
1. 排序效率问题
2. 文件编码问题
3. 内存使用问题

通过调试解决了这些问题。
''')

stu_doc.add_heading('心得体会', level=1)
stu_doc.add_paragraph('''
通过实训学到了很多：
1. 实践很重要
2. 需要多练习
3. 要善于解决问题
4. 代码规范很重要
''')

stu_doc.save('student_report.docx')

print("测试文档创建成功！")